"""Gemma (via Ollama) extracts promises; code validates every claim.

The model never decides pass/fail. Any claim that fails schema or grounding
checks is routed to human_review, never repaired.

Usage:  python extract_claims.py samples/runway-model gemma3:1b
        python extract_claims.py --selftest
"""
import json
import os
import re
import sys
import urllib.request
from pathlib import Path

CHECK_TYPES = ["spreadsheet_total", "human_review"]

SCHEMA = {
    "type": "object",
    "properties": {"claims": {"type": "array", "items": {
        "type": "object",
        "properties": {
            "claim_id": {"type": "string"},
            "claim_text": {"type": "string"},
            "source_file": {"type": "string"},
            "check_type": {"type": "string", "enum": CHECK_TYPES},
            "target": {"type": "string"},
            "expected_value": {"type": ["number", "null"]},
            "required_evidence": {"type": "string"},
        },
        "required": ["claim_id", "claim_text", "source_file", "check_type",
                     "target", "expected_value", "required_evidence"],
    }}},
    "required": ["claims"],
}

SYSTEM = """You extract testable promises from digital product documentation.
Return valid JSON only, matching the schema.

For each claim include: claim_id, claim_text (copied EXACTLY from the source),
source_file, check_type, target, expected_value, required_evidence.

check_type:
- spreadsheet_total: a sentence promises that a named total in the spreadsheet
  shows a specific number for the sample data. target must be EXACTLY one of the
  labels listed after "WORKBOOK TOTAL LABELS". expected_value is that number.
  claim_text is the whole sentence that makes the promise.
- human_review: any other concrete promise (quantities, time estimates,
  included items) that the supported checks cannot verify.

Rules:
1. Do not decide whether a claim is true.
2. Do not invent filenames, URLs, labels, numbers or requirements.
3. claim_text must be copied verbatim from the source.
4. Use an empty target and null expected_value when the source gives none.
5. Ignore promotional language such as "amazing" or "easy to use".
6. Do not extract file mentions or links; other code handles those.

Example input:
FILE: README.md
After loading the demo data, net profit comes to 2,310 for the quarter.

WORKBOOK TOTAL LABELS: Total Net Profit; Total Expenses

Example output:
{"claims": [{"claim_id": "c1",
 "claim_text": "After loading the demo data, net profit comes to 2,310 for the quarter.",
 "source_file": "README.md", "check_type": "spreadsheet_total",
 "target": "Total Net Profit", "expected_value": 2310,
 "required_evidence": "value of the Total Net Profit cell"}]}"""


def ollama_host():
    h = os.environ.get("OLLAMA_HOST", "http://127.0.0.1:11434")
    return h if h.startswith("http") else "http://" + h


def preflight(model):
    """Returns an error message if Ollama or the model is unavailable, else None."""
    try:
        with urllib.request.urlopen(ollama_host() + "/api/tags", timeout=10) as r:
            names = [m["name"] for m in json.loads(r.read()).get("models", [])]
    except Exception as e:
        return (f"Cannot reach Ollama at {ollama_host()} ({e}).\n"
                "Install it from ollama.com, start the Ollama app (or run `ollama serve` "
                "in another terminal), then run `ollama list` to check it works.")
    if model not in names and f"{model}:latest" not in names:
        return (f"Ollama is running but model '{model}' is not installed.\n"
                f"Installed models: {', '.join(names) or 'none'}.\n"
                f"Run `ollama pull {model}` or set OLLAMA_MODEL to one of the names above.")
    return None


def call_ollama(model, system, user, host=None):
    host = host or ollama_host()
    body = {"model": model, "stream": False, "format": SCHEMA,
            "options": {"temperature": 0},
            "messages": [{"role": "system", "content": system},
                         {"role": "user", "content": user}]}
    req = urllib.request.Request(host + "/api/chat", json.dumps(body).encode(),
                                 {"Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=300) as r:
        return json.loads(r.read())["message"]["content"]


def norm(s):
    return re.sub(r"\s+", " ", re.sub(r"[`*_]", "", s)).strip().lower()


def workbook_labels(root):
    """Total labels read from the .xlsx by plain code, given to the model as context."""
    from openpyxl import load_workbook
    out = []
    for x in sorted(Path(root).rglob("*.xlsx")):
        ws = load_workbook(x).active
        for row in range(2, ws.max_row + 1):
            v = ws.cell(row, 1).value
            if isinstance(v, str) and v.lower().startswith("total "):
                out.append(v.strip())
    return out


def validate(raw, sources, labels=None):
    """Returns (valid_claims, human_review). Never repairs a claim."""
    valid, review = [], []
    try:
        claims = json.loads(raw)["claims"]
        assert isinstance(claims, list)
    except Exception:
        return [], [{"reason": "invalid_json_or_schema", "raw": raw[:200]}]
    for c in claims:
        try:
            for k in ("claim_id", "claim_text", "source_file", "check_type", "target"):
                assert isinstance(c[k], str), f"missing_{k}"
            assert c["check_type"] in CHECK_TYPES, "unsupported_check_type"
            src = sources.get(c["source_file"])
            assert src is not None, "unknown_source_file"
            assert norm(c["claim_text"]) in norm(src), "claim_text_not_in_source"
            if c["check_type"] == "spreadsheet_total":
                v = c.get("expected_value")
                assert isinstance(v, (int, float)) and not isinstance(v, bool), "no_expected_value"
                shown = {str(int(v)), f"{int(v):,}", str(v)}
                assert any(s in c["claim_text"] for s in shown), "value_not_in_claim_text"
                if labels:
                    assert c["target"] in labels, "target_not_a_workbook_label"
                word = re.sub(r"^total\s+", "", c["target"].strip().lower())
                assert word and word in c["claim_text"].lower(), "target_not_in_claim_text"
            valid.append(c)
        except (AssertionError, KeyError, TypeError) as e:
            review.append({"reason": str(e) or "malformed", "claim": c})
    return valid, review


def extract(product_dir, model_fn):
    root = Path(product_dir)
    sources = {str(p.relative_to(root)): p.read_text(errors="ignore")
               for p in sorted(root.rglob("*")) if p.suffix.lower() in (".md", ".txt")}
    labels = workbook_labels(root)
    user = "\n\n".join(f"FILE: {n}\n{t}" for n, t in sources.items())
    if labels:
        user += "\n\nWORKBOOK TOTAL LABELS: " + "; ".join(labels)
    valid, review = validate(model_fn(SYSTEM, user), sources, labels)
    return [c for c in valid if c["check_type"] == "spreadsheet_total"], review


def selftest():
    src = {"README.md": "Revenue comes to 1,640 in total. Setup takes five minutes."}
    ok = {"claim_id": "c1", "claim_text": "Revenue comes to 1,640 in total.",
          "source_file": "README.md", "check_type": "spreadsheet_total",
          "target": "Total Revenue", "expected_value": 1640, "required_evidence": "sheet"}
    cases = {
        "valid": (ok, True),
        "invented_value": ({**ok, "expected_value": 1700}, False),
        "claim_not_in_source": ({**ok, "claim_text": "Revenue comes to 9 in total."}, False),
        "target_not_grounded": ({**ok, "target": "Total Profit"}, False),
        "unsupported_check_type": ({**ok, "check_type": "broken_link"}, False),
        "unknown_source": ({**ok, "source_file": "OTHER.md"}, False),
    }
    for name, (claim, expect_ok) in cases.items():
        v, r = validate(json.dumps({"claims": [claim]}), src)
        assert bool(v) == expect_ok, name
    v, r = validate(json.dumps({"claims": [ok]}), src, labels=["Total Cost"])
    assert not v and r[0]["reason"] == "target_not_a_workbook_label"
    v, r = validate("not json at all", src)
    assert not v and r[0]["reason"] == "invalid_json_or_schema"
    print(f"selftest passed ({len(cases) + 2} cases)")


if __name__ == "__main__":
    if "--selftest" in sys.argv:
        selftest()
    else:
        model = sys.argv[2] if len(sys.argv) > 2 else os.environ.get("OLLAMA_MODEL")
        if not model or len(sys.argv) < 2:
            sys.exit("Usage: python extract_claims.py <product_dir> <model tag>")
        problem = preflight(model)
        if problem:
            sys.exit(problem)
        claims, review = extract(sys.argv[1], lambda s, u: call_ollama(model, s, u))
        print(json.dumps({"claims": claims, "human_review": review}, indent=2))