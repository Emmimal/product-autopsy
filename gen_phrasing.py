"""Free-phrasing hold-out set: can a model catch promised totals the regex cannot parse?

20 mini products. Each has a correct workbook (no structural defect) and ONE sentence
promising the Total Revenue in natural wording. Ten templates, each used twice:
once with a WRONG promised number (defect) and once with the CORRECT number (control).
None of the sentences matches the deterministic regex, so the deterministic baseline
should catch 0 of 10. Whatever a model adds is measurable against that.

Written to samples_phrasing/ and ground_truth_phrasing.json (separate from the main set).
"""
import json
import random
import shutil
from pathlib import Path

from openpyxl import Workbook

HERE = Path(__file__).parent

TEMPLATES = [
    "Load the sample rows and revenue comes to {v} in total.",
    "With the demo data in place, you should see total revenue of {v}.",
    "After you paste the example rows, the revenue figure at the bottom reads {v}.",
    "The starter data adds up to {v} in revenue.",
    "Total revenue for the sample period: {v}.",
    "Expect {v} in revenue once the sample months are loaded.",
    "Run the sample data and the revenue line should land on {v}.",
    "Across the ten sample months, revenue totals {v}.",
    "You will get {v} as the revenue total from the included rows.",
    "The sample revenue column sums to {v}.",
]
NAMES = ["tutor-tracker", "candle-shop", "photo-booking", "bake-sale", "plant-seller",
         "print-farm", "tour-guide", "zine-shop", "pottery-log", "repair-desk"]


NOISY_README = """# {title}

{title} helps you keep a small business on track. Everything lives in one workbook.

## What's included
- The workbook: `tracker.xlsx`
- Demo rows: `example-data.csv`

## Setup
Open `tracker.xlsx` first and enter your starting balance.
{promise}
Your own rows replace the demo rows, and every total updates automatically.

## Links
- [Quick guide](guide.md)
- [Troubleshooting tips](guide.md#troubleshooting)
- [Jump to setup](#setup)
"""


def build(noisy=False):
    OUT = HERE / ("samples_phrasing_noisy" if noisy else "samples_phrasing")
    truth_name = "ground_truth_phrasing_noisy.json" if noisy else "ground_truth_phrasing.json"
    if OUT.exists():
        shutil.rmtree(OUT)
    truth = []
    n = 0
    for t, tmpl in enumerate(TEMPLATES):
        for wrong in (True, False):
            pid = f"{NAMES[t]}-{'wrong' if wrong else 'right'}"
            root = OUT / pid
            root.mkdir(parents=True)
            rng = random.Random(5000 + n)
            rev = [rng.randint(80, 220) for _ in range(10)]
            cost = [rng.randint(40, 120) for _ in range(10)]
            wb = Workbook()
            ws = wb.active
            ws.append(["Month", "Revenue", "Cost"])
            for m, (r, c) in enumerate(zip(rev, cost), start=1):
                ws.append([f"Month {m}", r, c])
            ws["A13"], ws["A14"] = "Total Revenue", "Total Cost"
            ws["B13"], ws["C14"] = "=SUM(B2:B11)", "=SUM(C2:C11)"
            wb.save(root / "tracker.xlsx")
            promised = sum(rev) + ((40 + 7 * t) * (1 if t % 2 == 0 else -1) if wrong else 0)
            promise = tmpl.format(v=f"{promised:,}")
            title = NAMES[t].replace("-", " ").title()
            if noisy:
                (root / "README.md").write_text(NOISY_README.format(title=title, promise=promise))
                (root / "example-data.csv").write_text(
                    "month,revenue,cost\n" + "\n".join(f"{m},{r},{c}" for m, (r, c)
                                                       in enumerate(zip(rev, cost), 1)))
                (root / "guide.md").write_text("# Guide\n\n## Setup\n\n## Troubleshooting\n")
            else:
                (root / "README.md").write_text(
                    f"# {title}\n\nA simple workbook, `tracker.xlsx`, for a small business.\n\n"
                    f"## Sample data\n{promise}\n")
            truth.append({"id": f"{pid}|sheet|total revenue", "product": pid,
                          "category": "spreadsheet", "defective": wrong,
                          "subtype": f"free_phrasing_t{t}_{'wrong' if wrong else 'right'}"})
            n += 1
    (HERE / truth_name).write_text(json.dumps(truth, indent=1))
    d = sum(x["defective"] for x in truth)
    print(f"{n} products, {d} wrong promises (defects), {n - d} correct promises (controls)")


if __name__ == "__main__":
    import sys
    build(noisy="--noisy" in sys.argv)